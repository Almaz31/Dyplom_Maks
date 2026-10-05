"""EM-DAT import, with explicit date and geography precision and unique event totals."""
import calendar
import csv
import hashlib
import io
import json
import math
import re
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

from .regions import REGIONS, region_code

TYPES = {'Flood':'Повінь','Storm':'Шторм','Extreme temperature':'Екстремальна температура','Drought':'Посуха',
 'Wildfire':'Природна пожежа','Air':'Авіаційна катастрофа','Road':'Дорожня аварія','Rail':'Залізнична аварія',
 'Water':'Водний транспорт','Explosion (Industrial)':'Промисловий вибух','Explosion (Miscellaneous)':'Інший вибух',
 'Fire (Industrial)':'Промислова пожежа','Fire (Miscellaneous)':'Інша пожежа','Collapse (Industrial)':'Промислове обвалення',
 'Chemical spill':'Викид хімічних речовин','Gas leak':'Витік газу','Miscellaneous accident (General)':'Інша аварія'}
GAUL = dict(zip(range(3148,3173), ['UA-71','UA-74','UA-77','UA-12','UA-14','UA-26','UA-63','UA-65','UA-68','UA-35','UA-43','UA-32','UA-46','UA-09','UA-48','UA-51','UA-53','UA-56','UA-59','UA-61','UA-05','UA-07','UA-21','UA-23','UA-18']))

def normal(text):
    return re.sub(r'[^a-zа-яіїєґ0-9]', '', str(text).casefold())

REGION_NAMES = {
 'UA-05':['vinnyts','vinnits'], 'UA-07':['volyn'], 'UA-09':['luhans','lougansk'],
 'UA-12':['dniprop','dneprop'], 'UA-14':['donetsk','donestk','donesk','donestka'],
 'UA-18':['zhytom','jitomir','joutomyr'], 'UA-21':['zakarp','transcarpath','zakarpat'],
 'UA-23':['zaporiz','zaporoz','zaporij'], 'UA-26':['ivanofrank','ivano rank'],
 'UA-32':['kyyiv','kyivoblast','kyivregion','kievregion','kievprovince'],
 'UA-35':['kirovo','kropyv'], 'UA-43':['crimea','crimee','crime','krym','krim','krimee'],
 'UA-46':['lviv'], 'UA-48':['mykol','nikola'], 'UA-51':['odes'], 'UA-53':['poltav'],
 'UA-56':['rivne','rovno'], 'UA-59':['sumy','sums','soumy'], 'UA-61':['ternop','temopil'],
 'UA-63':['khark','harkov'], 'UA-65':['kherson'], 'UA-68':['khmel','khmeln'],
 'UA-71':['cherkas','tcherkass'], 'UA-74':['chernih','tchernig'], 'UA-77':['cherniv','chernov','chemiv'],
 'UA-30':['capital kyiv','kyiv city','kiev city'], 'UA-40':['sevastopol'],
}

def inferred_regions(text):
    compact=normal(text)
    result={code for code,aliases in REGION_NAMES.items() if any(normal(a) in compact for a in aliases)}
    for code,name in REGIONS.items():
        if normal(name) in compact: result.add(code)
    return result

def numeric(value, factor=1):
    if value in [None,'','NA','NaN']: return None
    result=float(value)*factor
    if not math.isfinite(result) or result<0: raise ValueError('Некоректні наслідки')
    return result

def date_bounds(year,month,day):
    y=int(year)
    if not 1900<=y<=2100: raise ValueError('Рік поза підтримуваним діапазоном')
    if month in [None,'']:
        return f'{y}-01-01',f'{y}-12-31','year'
    m=int(month)
    if day in [None,'']:
        return date(y,m,1).isoformat(),date(y,m,calendar.monthrange(y,m)[1]).isoformat(),'month'
    d=date(y,m,int(day)).isoformat()
    return d,d,'day'

def parse_emdat(content):
    workbook=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    try:
        sheet=workbook['EM-DAT Data'] if 'EM-DAT Data' in workbook.sheetnames else workbook.active
        sheet.reset_dimensions()  # EM-DAT exports can incorrectly advertise a 1-cell dimension.
        iterator=sheet.iter_rows(values_only=True)
        headers=next(iterator)
        required={'DisNo.','ISO','Disaster Type','Start Year'}
        if not required.issubset(headers): raise ValueError('Очікується Excel EM-DAT з полями DisNo., ISO, Disaster Type, Start Year')
        events,errors,foreign,seen=[],[],0,set()
        for index,values in enumerate(iterator,2):
            if index>100002: raise ValueError('Завелика кількість записів')
            if not any(v not in [None,''] for v in values):continue
            row=dict(zip(headers,values))
            if str(row.get('ISO','')).strip()!='UKR': foreign+=1;continue
            try:
                ident=str(row['DisNo.']).strip()
                if not ident or ident in seen: raise ValueError('Порожній або повторний DisNo.')
                start_low,start_high,precision=date_bounds(row['Start Year'],row.get('Start Month'),row.get('Start Day'))
                if row.get('End Year') not in [None,'']:
                    end_low,end_high,_=date_bounds(row['End Year'],row.get('End Month'),row.get('End Day'))
                else:end_low,end_high=start_low,start_high
                if end_high<start_low:raise ValueError('Кінець події раніше початку')
                regions=set();basis=[]
                gadm=row.get('GADM Admin Units')
                if gadm:
                    for unit in json.loads(gadm):
                        names=' '.join(str(unit.get(k,'')) for k in ['name_1','adm1_name'])
                        regions.update(inferred_regions(names))
                    if regions:basis.append('GADM Admin Units')
                if row.get('Admin Units'):
                    for unit in json.loads(row['Admin Units']):
                        code=GAUL.get(unit.get('adm1_code'))
                        if code:regions.add(code)
                    if regions:basis.append('Admin Units / GAUL')
                location=str(row.get('Location') or '')
                # Add explicit region names in Location, but never guess an ambiguous village.
                extra=inferred_regions(location)-regions
                if extra:regions.update(extra);basis.append('Назва області в Location')
                lat=row.get('Latitude');lon=row.get('Longitude')
                if lat not in [None,''] and lon not in [None,'']:
                    lat,lon=float(lat),float(lon)
                    if not math.isfinite(lat) or not math.isfinite(lon) or not (-90<=lat<=90 and -180<=lon<=180):raise ValueError('Некоректні координати')
                else:lat=lon=None
                event_type=str(row['Disaster Type']).strip()
                events.append({'id':ident,'type':event_type,'type_label':TYPES.get(event_type,event_type),
                    'group':str(row.get('Disaster Group') or ''),'name':str(row.get('Event Name') or ''),
                    'location':location,'regions':sorted(regions),'region_basis':'; '.join(basis) or 'Область не визначено',
                    'latitude':lat,'longitude':lon,'start':start_low,'end':end_high,'start_high':start_high,'end_low':end_low,
                    'date_precision':precision,'year':int(row['Start Year']),
                    'deaths':numeric(row.get('Total Deaths')),'affected':numeric(row.get('Total Affected')),
                    'damage_usd':numeric(row.get("Total Damage ('000 US$)"),1000),
                    'damage_adjusted_usd':numeric(row.get("Total Damage, Adjusted ('000 US$)"),1000),
                    'last_update':str(row.get('Last Update') or ''),'source':'EM-DAT / CRED · власний Excel'})
                seen.add(ident)
            except (ValueError,TypeError,KeyError) as exc:
                errors.append({'row':index,'reason':str(exc)})
        if not events:raise ValueError('У файлі немає коректних записів для України')
        report={'count':len(events),'skipped_foreign':foreign,'errors':errors,
            'start':min(e['start'] for e in events),'end':max(e['end'] for e in events),
            'region_count':len({r for e in events for r in e['regions']}),
            'unlocated':sum(not e['regions'] and e['latitude'] is None for e in events),
            'incomplete_dates':sum(e['date_precision']!='day' for e in events),
            'types':dict(Counter(e['type'] for e in events))}
        return events,report
    finally:workbook.close()

class CommitConfig(BaseModel):
    name:str='Катастрофи України · EM-DAT'
    skip_invalid:bool=False

def create_router(connect,data):
    router=APIRouter(prefix='/api')
    uploads=data/'disaster_uploads';uploads.mkdir(exist_ok=True)
    with connect() as con:
        con.executescript('''CREATE TABLE IF NOT EXISTS disaster_datasets(id TEXT PRIMARY KEY,name TEXT,filename TEXT,created TEXT,report TEXT);
        CREATE TABLE IF NOT EXISTS disasters(dataset_id TEXT,id TEXT,start TEXT,end TEXT,type TEXT,payload TEXT,PRIMARY KEY(dataset_id,id));
        CREATE INDEX IF NOT EXISTS disasters_dates ON disasters(dataset_id,start,end);''')

    def load_upload(token):
        try:token=str(uuid.UUID(token))
        except ValueError:raise HTTPException(404,'Імпорт не знайдено')
        p=uploads/(token+'.json')
        if not p.exists():raise HTTPException(404,'Імпорт не знайдено')
        return json.loads(p.read_text(encoding='utf-8'))

    @router.post('/disaster-uploads')
    async def preview(file:UploadFile=File(...)):
        if not (file.filename or '').lower().endswith('.xlsx'):raise HTTPException(400,'Завантажте Excel .xlsx із EM-DAT')
        content=await file.read(20*1024*1024+1);await file.close()
        if len(content)>20*1024*1024:raise HTTPException(413,'Максимум 20 МБ')
        try:events,report=parse_emdat(content)
        except Exception as exc:raise HTTPException(400,f'Не вдалося прочитати EM-DAT: {exc}')
        token=str(uuid.uuid4())
        (uploads/(token+'.xlsx')).write_bytes(content)
        payload={'filename':file.filename,'events':events,'report':report,'sha256':hashlib.sha256(content).hexdigest()}
        (uploads/(token+'.json')).write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
        return {'token':token,'report':report,'preview':events[:5],'filename':file.filename}

    @router.post('/disaster-uploads/{token}/commit')
    def commit(token:str,config:CommitConfig):
        payload=load_upload(token)
        if payload['report']['errors'] and not config.skip_invalid:raise HTTPException(400,'Підтвердьте пропуск помилкових записів')
        with connect() as con:
            # Same original is idempotent: do not create duplicate catalogs accidentally.
            existing=con.execute('SELECT id,report FROM disaster_datasets').fetchall()
            for row in existing:
                if json.loads(row['report']).get('sha256')==payload['sha256']:return {'id':row['id'],'existing':True}
            ident=str(uuid.uuid4());report={**payload['report'],'sha256':payload['sha256'],'filename':payload['filename']}
            con.execute('INSERT INTO disaster_datasets VALUES(?,?,?,?,?)',(ident,config.name.strip()[:160] or 'EM-DAT',token,datetime.now().isoformat(),json.dumps(report,ensure_ascii=False)))
            con.executemany('INSERT INTO disasters VALUES(?,?,?,?,?,?)',[(ident,e['id'],e['start'],e['end'],e['type'],json.dumps(e,ensure_ascii=False)) for e in payload['events']])
        return {'id':ident}

    @router.get('/disaster-datasets')
    def catalog():
        with connect() as con:rows=con.execute('SELECT * FROM disaster_datasets ORDER BY created DESC').fetchall()
        return [{**dict(r),'report':json.loads(r['report'])} for r in rows]

    def events_for(ident,start,end,region,event_type,group):
        try:
            if date.fromisoformat(start)>date.fromisoformat(end):raise ValueError('Початок після кінця')
        except ValueError:raise HTTPException(400,'Некоректний період')
        if region!='all' and region not in REGIONS:raise HTTPException(400,'Невідома область')
        with connect() as con:
            ds=con.execute('SELECT * FROM disaster_datasets WHERE id=?',(ident,)).fetchone()
            if not ds:raise HTTPException(404,'Каталог не знайдено')
            events=[json.loads(r['payload']) for r in con.execute('SELECT payload FROM disasters WHERE dataset_id=? AND start<=? AND end>=? ORDER BY start',(ident,end,start))]
        events=[e for e in events if (region=='all' or region in e['regions']) and (not event_type or e['type']==event_type) and (not group or e['group']==group)]
        return events,dict(ds)

    @router.get('/disaster-datasets/{ident}/analysis')
    def analyze(ident:str,start:str,end:str,region:str='all',event_type:str='',group:str=''):
        events,ds=events_for(ident,start,end,region,event_type,group)
        count=len(events);by_region={}
        for code in REGIONS:
            involved=[e for e in events if code in e['regions']]
            by_region[code]={'value':len(involved),'coverage':100,'event_ids':[e['id'] for e in involved]}
        def impact(key):
            known=[e[key] for e in events if e[key] is not None]
            return {'value':sum(known) if known else None,'known':len(known),'unknown':count-len(known)}
        # Chart counts start occurrences, not every year touched by a long event.
        series=[]
        for year in range(int(start[:4]),int(end[:4])+1):
            row=[e for e in events if e['year']==year]
            deaths=[e['deaths'] for e in row if e['deaths'] is not None]
            series.append({'year':str(year),'count':len(row),'deaths':sum(deaths) if deaths else None})
        return {'events':events,'map':by_region,'series':series,
            'types':[{'type':k,'label':TYPES.get(k,k),'count':v} for k,v in Counter(e['type'] for e in events).most_common()],
            'stats':{'count':count,'deaths':impact('deaths'),'affected':impact('affected'),'damage':impact('damage_usd'),
                'located':sum(bool(e['regions']) for e in events),'unlocated':sum(not e['regions'] for e in events),
                'incomplete_dates':sum(e['date_precision']!='day' for e in events)},'source':'EM-DAT / CRED','report':json.loads(ds['report'])}

    @router.get('/disaster-datasets/{ident}/export')
    def export(ident:str,start:str,end:str,region:str='all',event_type:str='',group:str=''):
        events,_=events_for(ident,start,end,region,event_type,group)
        out=io.StringIO();fields=['id','type','group','start','end','date_precision','location','regions','region_basis','deaths','affected','damage_usd','damage_adjusted_usd']
        writer=csv.DictWriter(out,fields,extrasaction='ignore');writer.writeheader()
        for event in events:
            row={**event,'regions':';'.join(event['regions'])}
            for key,value in row.items():
                if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):row[key]="'"+value
            writer.writerow(row)
        return Response(out.getvalue().encode('utf-8-sig'),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="disasters-selection.csv"'})

    @router.get('/disaster-datasets/{ident}/original')
    def original(ident:str):
        with connect() as con:row=con.execute('SELECT filename FROM disaster_datasets WHERE id=?',(ident,)).fetchone()
        if not row:raise HTTPException(404,'Каталог не знайдено')
        return FileResponse(uploads/(row['filename']+'.xlsx'),filename='emdat-original.xlsx')

    return router
