"""Analytical views over existing immutable snapshots and imported SQLite datasets."""
import calendar
import csv
import gzip
import io
import json
import statistics
from datetime import date, datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field, model_validator
from openpyxl import Workbook

from .analytics_stats import (METRICS, SEASONS, aggregate, climate, descriptive, finite, forecast,
                              month_index, month_label, regression, relationships)
from .regions import REGIONS
from .city_labels import city_samples

ROOT = Path(__file__).resolve().parents[1]
UNITS = dict(temperature='°C', precipitation='мм', humidity='%', wind='м/с')
LABELS = dict(temperature='Температура', precipitation='Опади', humidity='Вологість', wind='Вітер')


@lru_cache(maxsize=1)
def nasa_snapshot(version):
    with gzip.open(ROOT/'samples/europe-weather-full.json.gz', 'rt', encoding='utf-8') as f:
        data = json.load(f)
    path = ROOT/'samples/europe-weather-observed.json.gz'
    if path.exists():
        with gzip.open(path, 'rt', encoding='utf-8') as f:
            extra = json.load(f)
        lookup = {p['id']: p for p in extra['points']}
        for p in data['points']:
            old = {(r['year'], r['month']) for r in p['monthly']}
            p['monthly'] += [r for r in lookup.get(p['id'], {}).get('monthly', []) if (r['year'], r['month']) not in old]
            p['monthly'].sort(key=lambda r: (r['year'], r['month']))
    return data


def nasa():
    paths = [ROOT/'samples/europe-weather-full.json.gz', ROOT/'samples/europe-weather-observed.json.gz']
    if not paths[0].exists():
        raise HTTPException(404, 'Локальний набір NASA POWER відсутній')
    return nasa_snapshot(tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size) for p in paths if p.exists()))


class Selection(BaseModel):
    source: str = Field(default='nasa', max_length=100)
    location: str = Field(default='country:UKR', max_length=100)
    start: str = Field(default='2000-01', pattern=r'^\d{4}-(0[1-9]|1[0-2])$')
    end: str = Field(default='2025-12', pattern=r'^\d{4}-(0[1-9]|1[0-2])$')
    metric: Literal['temperature', 'precipitation', 'humidity', 'wind'] = 'temperature'
    baseline_start: int = Field(default=1991, ge=1900, le=2090)
    baseline_end: int = Field(default=2020, ge=1900, le=2100)
    level: Literal['monthly', 'seasonal', 'yearly'] = 'monthly'
    rolling: int = Field(default=12, ge=2, le=60)
    season: Literal['all', 'Зима', 'Весна', 'Літо', 'Осінь'] = 'all'
    variables: list[Literal['temperature', 'precipitation', 'humidity', 'wind']] = Field(default_factory=lambda: list(METRICS), min_length=2, max_length=4)
    transform: Literal['raw', 'deseasonalized'] = 'deseasonalized'

    @model_validator(mode='after')
    def intervals(self):
        if self.start > self.end or self.baseline_start > self.baseline_end or self.baseline_end-self.baseline_start < 9:
            raise ValueError('Перевірте період; базовий інтервал має містити щонайменше 10 років')
        if not 1900 <= int(self.start[:4]) <= int(self.end[:4]) <= 2100:
            raise ValueError('Підтримуються роки 1900–2100')
        if len(set(self.variables)) != len(self.variables):
            raise ValueError('Показники не повинні повторюватися')
        return self


class Comparison(Selection):
    locations: list[str] = Field(min_length=2, max_length=4)


class ForecastSelection(Selection):
    horizon: int = Field(default=12, ge=1, le=36)
    test_months: int = Field(default=24, ge=12, le=60)
    method: Literal['naive', 'trend'] = 'trend'


class ExportSelection(ForecastSelection):
    locations: list[str] = Field(default_factory=list, max_length=4)


class HazardSelection(BaseModel):
    catalog: str
    start: date
    end: date
    event_type: str = ''
    frequency_weight: float = Field(default=.7, ge=0, le=1)
    severity_weight: float = Field(default=.3, ge=0, le=1)

    @model_validator(mode='after')
    def valid(self):
        if self.start > self.end or self.frequency_weight+self.severity_weight <= 0:
            raise ValueError('Перевірте дати й ваги індикатора')
        return self


def hazard_indicator(events, start, end, fw=.7, sw=.3):
    """Relative historical index, NOT population-normalized risk or a forecast."""
    years = ((end-start).days+1)/365.2425
    items = []
    for code, name in REGIONS.items():
        members = [e for e in events if code in e['regions']]
        known = [e['deaths'] for e in members if finite(e.get('deaths'))]
        # Multi-region deaths are not allocated to a region: they describe linked events.
        items.append(dict(code=code, name=name, count=len(members), frequency=len(members)/years,
                          severity=statistics.mean(known) if known else None, known=len(known),
                          coverage=len(known)/len(members) if members else 0))
    # Zero-event regions have zero frequency, unknown severity, and no composite score.
    # Never reward missing severity by treating it as zero.
    def norm(key):
        values = [r[key] for r in items if r[key] is not None]
        lo, hi = min(values), max(values)
        return {r['code']: (100*(r[key]-lo)/(hi-lo) if hi > lo else 0.) if r[key] is not None else None for r in items}
    frequency = norm('frequency')
    severity = norm('severity') if any(r['severity'] is not None for r in items) else {r['code']: None for r in items}
    for r in items:
        r['frequency_score'], r['severity_score'] = frequency[r['code']], severity[r['code']]
        r['score'] = (fw*r['frequency_score']+sw*r['severity_score'])/(fw+sw) if r['count'] and r['severity_score'] is not None else None
        if sw == 0 and r['count']:
            r['score'] = r['frequency_score']
    return items


def create_analytics_router(connect):
    router = APIRouter(prefix='/api/analytics')

    def source_info():
        sources = []
        if (ROOT/'samples/europe-weather-full.json.gz').exists():
            data = nasa()
            points = data['points']
            dates = [f"{r['year']:04d}-{r['month']:02d}" for r in points[0]['monthly']]
            countries = sorted({(p['iso'], p['country']) for p in points}, key=lambda x: x[1])
            provenance = data.get('provenance', [])
            retrieved = max((p.get('retrieved_at', '') for p in provenance), default='')
            sources.append(dict(id='nasa', name='NASA POWER · Європа', source='NASA POWER / MERRA-2',
                                start=dates[0], end=dates[-1], kind='grid', retrieved_at=retrieved,
                                variables=list(METRICS), locations=[dict(id='country:'+iso, name=name) for iso, name in countries]+
                                [dict(id=ident, name=f"Поблизу міста {city['name']} · {city['point']['country']}")
                                 for ident, city in city_samples(points).items()]))
        with connect() as con:
            datasets = con.execute('SELECT * FROM datasets ORDER BY created DESC').fetchall()
        for d in datasets:
            report = json.loads(d['report'])
            sources.append(dict(id='dataset:'+d['id'], name=d['name'], source=d['source'], kind=d['spatial_kind'],
                                start=report['start'][:7], end=report['end'][:7], retrieved_at=d['created'],
                                variables=report['variables'], locations=[dict(id=r, name=REGIONS.get(r, r)) for r in report['regions']]))
        return sources

    def load(source, location):
        if source == 'nasa':
            data = nasa()
            city=city_samples(data['points']).get(location)
            sample_id=city['point']['id'] if city else location
            members = [p for p in data['points'] if (p['iso'] == location[8:] if location.startswith('country:') else p['id'] == sample_id)]
            if not members:
                raise HTTPException(400, 'Територію не знайдено в цьому джерелі')
            weights = [p.get('weight', 1.) if location.startswith('country:') else 1. for p in members]
            if not sum(weights):
                weights = [1.] * len(members)
            lookups = [{(r['year'], r['month']): r for r in p['monthly']} for p in members]
            keys = sorted(set().union(*(v.keys() for v in lookups)))
            rows = []
            for y, m in keys:
                item = dict(date=f'{y:04d}-{m:02d}')
                for metric in METRICS:
                    values = [(p.get((y, m), {}).get(metric), w) for p, w in zip(lookups, weights) if w > 0]
                    item[metric] = sum(v*w for v, w in values)/sum(w for v, w in values) if all(finite(v) for v, w in values) else None
                rows.append(item)
            label = f"Поблизу міста {city['name']}" if city else members[0]['country'] if location.startswith('country:') else f"{members[0]['country']} · вибрана територія"
            return rows, dict(source='NASA POWER / MERRA-2', location=label, spatial_kind='Індекс сітки країни' if location.startswith('country:') else 'Точка сітки',
                              source_url='https://power.larc.nasa.gov/', retrieved_at=max((p.get('retrieved_at', '') for p in data.get('provenance', [])), default=''),
                              note=(f"Використано найближчу доступну точку реаналізу: {city['distance_km']} км від міста. Це не метеостанція міста. " if city else '')+'Місячні оцінки реаналізу; індекс країни зважено за cos(широти), це не точний середній показник адміністративного полігона. Додаткові місяці — з окремого фактичного знімка; не прогноз.')
        if not source.startswith('dataset:') or location not in REGIONS:
            raise HTTPException(400, 'Невідоме джерело або територія')
        ident = source[8:]
        with connect() as con:
            ds = con.execute('SELECT * FROM datasets WHERE id=?', (ident,)).fetchone()
            if not ds:
                raise HTTPException(404, 'Набір не знайдено')
            report = json.loads(ds['report'])
            if location not in report['regions']:
                raise HTTPException(400, 'Територія відсутня у наборі')
            # SQL aggregates the stored daily rows; only monthly results reach the browser.
            expressions = ','.join(f'{"SUM" if k == "precipitation" else "AVG"}({k}) AS {k},COUNT({k}) AS {k}_n' for k in METRICS)
            monthly = con.execute(f"SELECT substr(date,1,7) AS month,{expressions} FROM observations WHERE dataset_id=? AND region=? GROUP BY month ORDER BY month", (ident, location)).fetchall()
        lookup = {r['month']: r for r in monthly}
        rows = []
        for ix in range(month_index(report['start']), month_index(report['end'])+1):
            label = month_label(ix)
            days = calendar.monthrange(ix//12, ix%12+1)[1]
            r = lookup.get(label)
            rows.append(dict(date=label, **{k: r[k] if r and r[k+'_n'] == days and finite(r[k]) else None for k in METRICS}))
        return rows, dict(source=ds['source'], location=REGIONS[location], spatial_kind='Точка біля міста' if ds['spatial_kind'] == 'point' else 'Регіональні дані',
                          retrieved_at=ds['created'], source_url='',
                          note='Аналізуються лише повні календарні місяці для кожного показника. Дата набору — дата імпорту, не підтвердження актуальності джерела.')

    @router.get('/sources')
    def sources():
        return dict(sources=source_info(), units=UNITS, variables=LABELS)

    def report(s):
        rows, meta = load(s.source, s.location)
        available_start, available_end = (rows[0]['date'], rows[-1]['date']) if rows else (None, None)
        # Explicit missing buckets prevent a selected 12-month window from appearing
        # 100% complete when the snapshot ends after only eight months.
        lookup = {r['date']: r for r in rows}
        first = min(s.start, available_start) if available_start else s.start
        last = max(s.end, available_end) if available_end else s.end
        rows = [lookup.get(month_label(ix), dict(date=month_label(ix), **{k: None for k in METRICS}))
                for ix in range(month_index(first), month_index(last)+1)]
        result = climate(rows, s.metric, s.start, s.end, s.baseline_start, s.baseline_end, s.level, s.rolling, s.season)
        selected = [r for r in rows if s.start <= r['date'] <= s.end and (s.season == 'all' or SEASONS[int(r['date'][5:7])] == s.season)]
        corr_rows = []
        # Remove each calendar month's mean over the selected sample, not a future dataset.
        means = {(m, k): statistics.mean(v) for m in range(1, 13) for k in s.variables
                 if (v := [r[k] for r in selected if int(r['date'][5:7]) == m and finite(r[k])])}
        for r in selected:
            item = dict(date=r['date'])
            for k in s.variables:
                item[k] = (r[k]-means[(int(r['date'][5:7]), k)] if s.transform == 'deseasonalized' else r[k]) if finite(r[k]) else None
            corr_rows.append(item)
        yearly_temperature = aggregate(rows, 'temperature', 'yearly')
        yearly_rain = aggregate(rows, 'precipitation', 'yearly')
        annual = [dict(year=r['year'], temperature=r['value'], precipitation=next((p['value'] for p in yearly_rain if p['year'] == r['year']), None)) for r in yearly_temperature if f"{r['year']}-01" >= s.start and f"{r['year']}-12" <= s.end]
        annual_fit = regression([(r['year'], r['temperature']) for r in annual])
        valid_annual = [r for r in annual if r['temperature'] is not None]
        change = valid_annual[-1]['temperature']-valid_annual[0]['temperature'] if len(valid_annual) >= 2 else None
        result.update(meta=meta, filters=s.model_dump(), unit=UNITS[s.metric],
                      correlations=relationships(corr_rows, s.variables), aligned=corr_rows,
                      descriptions={k: descriptive([r[k] for r in selected]) for k in s.variables},
                      annual=annual, annual_trend=annual_fit, annual_change=change,
                      quality=dict(months=len(selected), complete=sum(finite(r[s.metric]) for r in selected),
                                   available_start=available_start, available_end=available_end))
        return result

    @router.post('/report')
    def analytical_report(s: Selection):
        return report(s)

    @router.post('/forecast')
    def forecast_report(s: ForecastSelection):
        rows, meta = load(s.source, s.location)
        selected = [r for r in rows if s.start <= r['date'] <= s.end]
        if not selected or selected[0]['date'] != s.start or selected[-1]['date'] != s.end:
            raise HTTPException(400, 'Вибраний період виходить за межі наявної історії')
        if s.season != 'all':
            raise HTTPException(400, 'Прогнозування потребує всіх сезонів; вимкніть сезонний фільтр')
        try:
            result = forecast(selected, s.metric, s.horizon, s.test_months, s.method)
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        return dict(**result, meta=meta, unit=UNITS[s.metric], filters=s.model_dump())

    @router.post('/compare')
    def compare(s: Comparison):
        if len(set(s.locations)) != len(s.locations):
            raise HTTPException(400, 'Виберіть різні території')
        locations = []
        for location in s.locations:
            rows, meta = load(s.source, location)
            lookup = {r['date']: r[s.metric] for r in rows if s.start <= r['date'] <= s.end and (s.season == 'all' or SEASONS[int(r['date'][5:7])] == s.season)}
            locations.append(dict(id=location, meta=meta, lookup=lookup))
        shared = sorted(set.intersection(*({k for k, v in p['lookup'].items() if finite(v)} for p in locations)))
        series = [dict(date=key, **{p['id']: p['lookup'][key] for p in locations}) for key in shared]
        for p in locations:
            p['stats'] = descriptive([p['lookup'][d] for d in shared])
            p['trend'] = regression([(month_index(d)/12, p['lookup'][d]) for d in shared])
            p.pop('lookup')
        return dict(locations=locations, rows=series, common_months=len(shared), unit=UNITS[s.metric], filters=s.model_dump())

    @router.post('/hazards')
    def hazards(s: HazardSelection):
        with connect() as con:
            ds = con.execute('SELECT * FROM disaster_datasets WHERE id=?', (s.catalog,)).fetchone()
            if not ds:
                raise HTTPException(404, 'Каталог не знайдено')
            rows = con.execute('SELECT payload FROM disasters WHERE dataset_id=? AND start BETWEEN ? AND ? ORDER BY start', (s.catalog, s.start.isoformat(), s.end.isoformat())).fetchall()
        available = json.loads(ds['report'])
        if s.start.isoformat() < available['start'] or s.end.isoformat() > available['end']:
            raise HTTPException(400, 'Виберіть період усередині часових меж каталогу')
        events = [json.loads(r['payload']) for r in rows]
        events = [e for e in events if e['group'] == 'Natural' and (not s.event_type or e['type'] == s.event_type)]
        items = hazard_indicator(events, s.start, s.end, s.frequency_weight, s.severity_weight)
        return dict(regions=items, events=events, count=len(events), unlocated=sum(not e['regions'] for e in events),
                    series=[dict(year=y, count=sum(e['year'] == y for e in events)) for y in range(s.start.year, s.end.year+1)],
                    weights=dict(frequency=s.frequency_weight, severity=s.severity_weight), source='EM-DAT / CRED',
                    note='Історичний індикатор небезпеки, не повний ризик і не прогноз. Частота = унікальні пов’язані події / тривалість у роках. Тяжкість = середнє відомих загиблих у пов’язаних подіях; ці числа не розподілено між областями. Обидві компоненти min–max 0–100 між областями вибірки; ваги користувача нормуються. За невідомої тяжкості спільний індекс відсутній. Експозиція й вразливість не вимірюються.')

    def tabular_export(kind, sections, metadata):
        """Multiple result tables with provenance; all content computed server-side."""
        labels = dict(date='Період', year='Рік', actual='Фактичне значення', predicted='Прогноз',
                      baseline='Сезонний наївний', error='Факт − прогноз', low='Нижня межа', high='Верхня межа',
                      calibration_n='Помилок калібрування', value='Значення', normal='Базове середнє',
                      anomaly='Абсолютна аномалія', precipitation_pct='Відхилення опадів, %',
                      z='z-оцінка', rolling='Ковзне середнє', yoy='Рік до року', baseline_count='Років бази',
                      temperature='Температура, °C', precipitation='Опади, мм', humidity='Вологість, %', wind='Вітер, м/с',
                      variable='Показник', unit='Одиниця', count='Кількість', mean='Середнє', median='Медіана', sd='σ',
                      min='Мінімум', max='Максимум', outliers='Викиди 1,5 IQR', model='Модель', n='Кількість пар',
                      slope='Нахил', intercept='Вільний член', pearson='Пірсон', spearman='Спірмен', r2='R²',
                      name='Територія', code='Код', frequency='Подій на рік', severity='Середнє загиблих у пов’язаних подіях',
                      known='Відомі наслідки', coverage='Частка відомих наслідків', score='Індикатор',
                      frequency_score='Нормована частота', severity_score='Нормована тяжкість', location='Територія',
                      x='Показник X', y='Показник Y', type='Тип', type_label='Назва типу', regions='Коди областей')
        def safe(v):
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False)
            return "'"+v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v
        metadata = {**metadata, 'exported_at': datetime.now(timezone.utc).isoformat()}
        if kind == 'csv':
            out = io.StringIO()
            writer = csv.writer(out)
            for key, value in metadata.items():
                writer.writerow(['# '+key, safe(value)])
            for name, rows in sections:
                writer.writerow(['# '+name])
                if not rows:
                    continue
                fields = list(rows[0])
                writer.writerow([labels.get(k, k) for k in fields])
                writer.writerows([[safe(row.get(k)) for k in fields] for row in rows])
            return Response(out.getvalue().encode('utf-8-sig'), media_type='text/csv', headers={'Content-Disposition': 'attachment; filename="obrii-results.csv"'})
        workbook = Workbook()
        workbook.active.title = 'Метадані'
        for key, value in metadata.items():
            workbook.active.append([key, safe(value)])
        for name, rows in sections:
            sheet = workbook.create_sheet(name[:31])
            if rows:
                keys = list(rows[0])
                sheet.append([labels.get(k, k) for k in keys])
                for row in rows:
                    sheet.append([safe(row.get(k)) for k in keys])
                sheet.freeze_panes = 'A2'
                sheet.auto_filter.ref = sheet.dimensions
        out = io.BytesIO()
        workbook.save(out)
        return Response(out.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="obrii-results.xlsx"'})

    @router.post('/hazards/export/{kind}')
    def export_hazards(s: HazardSelection, kind: Literal['csv', 'xlsx']):
        result = hazards(s)
        return tabular_export(kind, [('Області', result['regions']), ('Події', result['events']), ('Роки', result['series'])],
                              dict(source=result['source'], filters=s.model_dump(mode='json'), methodology=result['note']))

    @router.post('/export/{kind}')
    def export(s: ExportSelection, kind: Literal['csv', 'xlsx'], view: Literal['climate', 'statistics', 'history', 'forecast', 'comparison'] = 'climate'):
        if view == 'forecast':
            result = forecast_report(s)
            return tabular_export(kind, [('Історія', result['history']), ('Перевірка', result['validation']), ('Прогноз', result['future']),
                                         ('Точність', [dict(model='Обрана модель', **result['accuracy']), dict(model='Сезонний наївний', **result['baseline_accuracy'])])],
                                  dict(**result['meta'], filters=s.model_dump(), units=result['unit'], training_end=result['training_end'], test_start=result['test_start'], test_end=result['test_end'], interval_note=result['interval_note']))
        if view == 'comparison':
            if not 2 <= len(s.locations) <= 4 or len(set(s.locations)) != len(s.locations):
                raise HTTPException(400, 'Для порівняння виберіть від двох до чотирьох різних територій')
            result = compare(Comparison(**s.model_dump()))
            return tabular_export(kind, [('Місяці', result['rows']), ('Показники', [dict(location=p['meta']['location'], **p['stats']) for p in result['locations']])],
                                  dict(filters=s.model_dump(), unit=result['unit'], locations=[p['meta'] for p in result['locations']], common_months=result['common_months']))
        result = report(s)
        if view == 'statistics':
            return tabular_export(kind, [('Описова статистика', [dict(variable=k, unit=UNITS[k], **v) for k, v in result['descriptions'].items()]),
                                         ('Кореляції', result['correlations']), ('Вирівняні місяці', result['aligned'])],
                                  dict(**result['meta'], filters=s.model_dump(), transform=s.transform,
                                       methodology_note='Пірсон і Спірмен за попарно сумісними місяцями; кореляція не доводить причинності.'))
        if view == 'history':
            return tabular_export(kind, [('Річні значення', result['annual']), ('Місячні аномалії', result['rows']),
                                         ('Річний тренд', [result['annual_trend']])],
                                  dict(**result['meta'], filters=s.model_dump(), units=UNITS, annual_change=result['annual_change'],
                                       methodology_note='Річна температура зважена за днями місяців; річні опади — сума 12 повних місяців.'))
        fields = [('date', 'Період'), ('value', LABELS[s.metric]+', '+result['unit']), ('normal', 'Базове середнє'),
                  ('anomaly', 'Відхилення'), ('precipitation_pct', 'Опади: відхилення, %'), ('z', 'z-оцінка'),
                  ('rolling', 'Ковзне середнє'), ('yoy', 'Різниця з попереднім роком'), ('baseline_count', 'Років бази')]
        meta = {**result['meta'], **s.model_dump(), 'exported_at': datetime.now(timezone.utc).isoformat(),
                'standard_30_year': result['baseline']['standard_30_year'], 'unit': result['unit']}
        def safe(v):
            return "'"+v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v
        if kind == 'csv':
            out = io.StringIO()
            writer = csv.writer(out)
            for k, v in meta.items():
                writer.writerow(['# '+k, safe(str(v))])
            writer.writerow([label for key, label in fields])
            writer.writerows([[safe(r.get(k)) for k, label in fields] for r in result['rows']])
            return Response(out.getvalue().encode('utf-8-sig'), media_type='text/csv', headers={'Content-Disposition': 'attachment; filename="obrii-analysis.csv"'})
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Аналітика'
        sheet.append([label for key, label in fields])
        for r in result['rows']:
            sheet.append([safe(r.get(k)) for k, label in fields])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        metadata = workbook.create_sheet('Метадані')
        for key, value in meta.items():
            metadata.append([key, safe(str(value))])
        stats = workbook.create_sheet('Статистика')
        for key, value in result['stats'].items():
            stats.append([key, value])
        correlations = workbook.create_sheet('Кореляції')
        keys = ['x', 'y', 'n', 'pearson', 'spearman', 'slope', 'intercept', 'r2']
        correlations.append(keys)
        for row in result['correlations']:
            correlations.append([row[k] for k in keys])
        out = io.BytesIO()
        workbook.save(out)
        return Response(out.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="obrii-analysis.xlsx"'})

    return router
