"""Build the editable diploma note. Requires python-docx (document tooling only)."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'docs/Пояснювальна записка.md'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2), Cm(2)
section.left_margin, section.right_margin = Cm(3), Cm(1.5)
normal = document.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.space_after = Pt(0)
for name in ('Title', 'Heading 1', 'Heading 2'):
    style = document.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(16 if name != 'Heading 2' else 14)
    style.font.bold = True
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(10)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
in_contents = False
for line in source.read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line.startswith('## ЗМІСТ'):
        document.add_page_break()
        document.add_heading('ЗМІСТ', level=1)
        in_contents = True
        continue
    if in_contents and line != '## ВСТУП':
        p = document.add_paragraph(line)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        continue
    if line == '## ВСТУП':
        in_contents = False
    if line.startswith('# '):
        p = document.add_heading(line[2:], level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('## '):
        title = line[3:]
        if title.startswith(('АНОТАЦІЯ', 'ВСТУП', '1.', '2.', 'ВИСНОВКИ', 'СПИСОК')):
            document.add_page_break()
        p = document.add_heading(title, level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('### '):
        document.add_heading(line[4:], level=2)
    elif line.startswith('- '):
        document.add_paragraph(line[2:], style='List Bullet')
    else:
        p = document.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if re.match(r'^\d+\. ', line):
            p.paragraph_format.first_line_indent = Cm(0)

target = source.with_suffix('.docx')
document.save(target)
print(f'Document created successfully; {len(document.paragraphs)} paragraphs')
