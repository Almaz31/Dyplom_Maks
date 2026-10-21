import io
import math
from datetime import date

import openpyxl
import pytest

from backend.analytics_stats import (aggregate, climate, descriptive, forecast, month_label,
                                     ranks, regression, relationships)
from backend.analytics_api import hazard_indicator
from tests.test_api import client, preview


def monthly(first=1981, years=40):
    return [dict(date=f'{y}-{m:02}', temperature=m+.1*(y-first), precipitation=2*m,
                 humidity=50., wind=2.) for y in range(first, first+years) for m in range(1, 13)]


def test_named_cities_resolve_existing_samples_without_coordinate_labels():
    sources=client.get('/api/analytics/sources').json()['sources']
    locations=next(source for source in sources if source['id']=='nasa')['locations']
    assert all('°' not in place['name'] for place in locations)
    assert any(place['id']=='city:lviv' and 'Львів' in place['name'] for place in locations)
    assert any(place['id']=='city:paris' and 'Париж' in place['name'] for place in locations)
    city=client.post('/api/analytics/report',json={'location':'city:lviv'}).json()
    grid=client.post('/api/analytics/report',json={'location':'50:24'}).json()
    assert city['rows']==grid['rows']
    assert city['meta']['location']=='Поблизу міста Львів'
    assert '18 км' in city['meta']['note']


def test_reference_statistics_and_tied_ranks():
    s = descriptive([1, 2, 3, 4, None, float('nan')])
    assert s['count'] == 4 and s['mean'] == s['median'] == 2.5
    assert s['sd'] == pytest.approx(math.sqrt(1.25))
    assert s['q1'] == 1.75 and s['q3'] == 3.25
    assert ranks([10, 10, 20, 5]) == [2.5, 2.5, 4, 1]
    assert descriptive([])['mean'] is None


def test_correlations_alignment_constants_and_small_samples():
    rows = [dict(a=x, b=2*x+1) for x in [1, 2, 3, 4]] + [dict(a=None, b=90)]
    r = relationships(rows, ['a', 'b'])[1]
    assert r['n'] == 4
    assert r['pearson'] == r['spearman'] == r['r2'] == 1
    assert r['slope'] == 2 and r['intercept'] == 1
    assert regression([(1, 2), (2, 4)])['slope'] is None
    assert regression([(1, 4), (2, 4), (3, 4)])['pearson'] is None
    assert regression([(1, 1), (1, 2), (1, 3)])['r2'] is None
    assert regression([(1, 6), (2, 4), (3, 2)])['pearson'] == -1


def test_climate_normal_rolling_yoy_and_missing_baseline():
    rows = monthly()
    r = climate(rows, 'temperature', '2011-01', '2011-12', 1981, 2010, window=3)
    assert r['baseline']['standard_30_year']
    first = r['rows'][0]
    assert first['normal'] == pytest.approx(2.45)
    assert first['anomaly'] == pytest.approx(1.55)
    assert first['yoy'] == pytest.approx(.1)
    assert first['rolling'] == pytest.approx((13.9+14.9+4)/3)
    assert first['z'] > 0
    rows[0]['temperature'] = None
    r = climate(rows, 'temperature', '2011-01', '2011-12', 1981, 2010)
    assert not r['baseline']['standard_30_year'] and r['rows'][0]['normal'] is None
    assert r['rows'][1]['normal'] is not None


def test_zero_normal_and_zero_variance_do_not_generate_infinity():
    rows = monthly()
    for r in rows:
        r['precipitation'] = 0
    r = climate(rows, 'precipitation', '2011-01', '2011-12', 1981, 2010)
    assert all(p['precipitation_pct'] is None and p['z'] is None for p in r['rows'])
    assert r['rows'][0]['anomaly'] == 0


def test_complete_year_and_cross_year_season_are_required():
    rows = monthly()
    annual = aggregate(rows, 'precipitation', 'yearly')
    assert annual[0]['value'] == sum(2*m for m in range(1, 13))
    winter = next(r for r in aggregate(rows, 'precipitation', 'seasonal') if r['year'] == 1982 and r['key'] == 'Зима')
    assert winter['value'] == 2*(12+1+2)
    assert aggregate(rows[1:], 'temperature', 'yearly')[0]['value'] is None
    assert climate(rows, 'precipitation', '2011-02', '2011-12', 1981, 2010, 'yearly')['rows'] == []


def test_forecast_fixed_origin_accuracy_and_no_future_leakage():
    rows = monthly(years=30)
    r = forecast(rows, 'temperature', 12, 24, 'trend')
    assert r['accuracy']['mae'] == pytest.approx(0, abs=1e-10)
    assert r['accuracy']['rmse'] == pytest.approx(0, abs=1e-10)
    assert r['baseline_accuracy']['mae'] > 0
    assert r['training_end'] < r['test_start']
    assert len(r['future']) == 12
    changed = [dict(p, temperature=p['temperature']+100 if p['date'] >= r['test_start'] else p['temperature']) for p in rows]
    after = forecast(changed, 'temperature', 12, 24, 'trend')
    assert [p['predicted'] for p in after['validation']] == [p['predicted'] for p in r['validation']]
    assert [p['calibration_n'] for p in after['future']] == [p['calibration_n'] for p in r['future']]
    for p in r['future']:
        assert p['low'] <= p['predicted'] <= p['high']
    with pytest.raises(ValueError):
        forecast(rows[:35], 'temperature')
    rows[60]['temperature'] = None
    with pytest.raises(ValueError):
        forecast(rows, 'temperature')


def test_hazard_scores_missing_severity_and_known_zeros():
    events = [dict(regions=['UA-46'], deaths=0), dict(regions=['UA-51'], deaths=10),
              dict(regions=['UA-51'], deaths=None), dict(regions=['UA-30'], deaths=None)]
    r = {p['code']: p for p in hazard_indicator(events, date(2020, 1, 1), date(2020, 12, 31))}
    assert r['UA-51']['score'] == 100
    assert r['UA-46']['score'] == 35
    assert r['UA-30']['score'] is None
    assert r['UA-51']['known'] == 1 and r['UA-51']['coverage'] == .5
    frequency_only = {p['code']: p for p in hazard_indicator(events, date(2020, 1, 1), date(2020, 12, 31), 1, 0)}
    assert frequency_only['UA-30']['score'] == 50


def test_real_snapshot_report_comparison_forecast_and_exports():
    sources = client.get('/api/analytics/sources').json()['sources']
    assert sources[0]['id'] == 'nasa'
    s = dict(source='nasa', location='country:UKR', start='2000-01', end='2025-12')
    response = client.post('/api/analytics/report', json=s)
    assert response.status_code == 200, response.text
    r = response.json()
    assert r['baseline']['standard_30_year']
    assert r['quality']['complete'] == 312
    assert len(r['correlations']) == 16
    assert r['meta']['source'] == 'NASA POWER / MERRA-2'
    comparison = client.post('/api/analytics/compare', json={**s, 'locations': ['country:UKR', 'country:POL']}).json()
    assert comparison['common_months'] == 312
    prediction = client.post('/api/analytics/forecast', json=s)
    assert prediction.status_code == 200, prediction.text
    assert prediction.json()['accuracy']['n'] == 24
    assert client.post('/api/analytics/forecast', json={**s, 'season': 'Зима'}).status_code == 400
    xlsx = client.post('/api/analytics/export/xlsx', json=s)
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx.content), read_only=True)
    assert workbook.sheetnames == ['Аналітика', 'Метадані', 'Статистика', 'Кореляції']
    assert workbook['Аналітика'].max_row == 313
    workbook.close()
    csv = client.post('/api/analytics/export/csv', json=s)
    assert '# source' in csv.content.decode('utf-8-sig')
    assert client.post('/api/analytics/report', json={**s, 'start': '2025-13'}).status_code == 422
    assert client.post('/api/analytics/report', json={**s, 'location': 'absent'}).status_code == 400


def test_partial_daily_month_is_not_silently_a_climate_total():
    p = preview('date,region,temperature,precipitation\n2024-01-01,UA-46,10,0\n2024-01-02,UA-46,20,1\n')
    ident = client.post('/api/uploads/'+p['token']+'/commit', json={'name': 'Неповний', 'mapping': p['mapping']}).json()['id']
    r = client.post('/api/analytics/report', json=dict(source='dataset:'+ident, location='UA-46', start='2024-01', end='2024-01')).json()
    assert r['rows'][0]['value'] is None
    assert r['quality']['complete'] == 0
    assert not r['baseline']['standard_30_year']


def test_analytics_reports_missing_end_of_snapshot_and_forecast_boundaries():
    s = dict(source='nasa', location='country:UKR', start='2026-01', end='2026-12')
    result = client.post('/api/analytics/report', json=s).json()
    assert result['quality']['months'] == 12
    assert result['quality']['complete'] == 8
    assert result['rows'][-1]['value'] is None
    assert client.post('/api/analytics/forecast', json=s).status_code == 400


def test_forecast_exports_are_derived_from_models_and_include_validation():
    s = dict(source='nasa', location='country:UKR', start='2000-01', end='2025-12', horizon=6)
    response = client.post('/api/analytics/export/xlsx?view=forecast', json=s)
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True)
    assert workbook['Прогноз'].max_row == 7
    assert workbook['Перевірка'].max_row == 25
    assert workbook['Точність'].max_row == 3
    workbook.close()
    for view, required in [('statistics', ['Описова статистика', 'Кореляції', 'Вирівняні місяці']),
                           ('history', ['Річні значення', 'Місячні аномалії', 'Річний тренд'])]:
        exported = client.post('/api/analytics/export/xlsx?view='+view, json=s)
        assert exported.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(exported.content), read_only=True)
        assert all(name in workbook.sheetnames for name in required)
        workbook.close()
        exported_csv = client.post('/api/analytics/export/csv?view='+view, json=s)
        assert required[0] in exported_csv.content.decode('utf-8-sig')
    comparison = client.post('/api/analytics/export/xlsx?view=comparison', json={**s, 'locations': ['country:UKR', 'country:POL']})
    assert comparison.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(comparison.content), read_only=True)
    assert workbook['Місяці'].max_row == 313
    workbook.close()


def test_network_weather_validation_rejects_duplicates_ranges_and_nonfinite():
    from fastapi import HTTPException
    from backend.world import validate_weather
    row = dict(latitude=50., longitude=30., daily=dict(time=['2024-01-01'], temperature_2m_mean=[0.],
               precipitation_sum=[0.], relative_humidity_2m_mean=[None], wind_speed_10m_mean=[1.]))
    assert validate_weather(row, date(2024, 1, 1), date(2024, 1, 1))['daily']['relative_humidity_2m_mean'] == [None]
    row['daily']['precipitation_sum'] = [-1]
    with pytest.raises(HTTPException):
        validate_weather(row, date(2024, 1, 1), date(2024, 1, 1))
    row['daily']['precipitation_sum'] = [float('nan')]
    with pytest.raises(HTTPException):
        validate_weather(row, date(2024, 1, 1), date(2024, 1, 1))
    row['daily']['time'] = ['2024-01-02']
    with pytest.raises(HTTPException):
        validate_weather(row, date(2024, 1, 1), date(2024, 1, 1))


def test_hazards_api_natural_only_uncertain_geography_and_exports():
    from tests.test_disasters import workbook_bytes
    book = openpyxl.load_workbook(io.BytesIO(workbook_bytes()))
    book.active.append(['2001-0009-UKR', 'UKR', 'Road', 'Technological', 2001, 6, 1, 'Lviv', '', 3, None, None])
    content = io.BytesIO()
    book.save(content)
    book.close()
    imported = client.post('/api/disaster-uploads', files={'file': ('hazards.xlsx', content.getvalue())}).json()
    ident = client.post('/api/disaster-uploads/'+imported['token']+'/commit', json={}).json()['id']
    s = dict(catalog=ident, start=imported['report']['start'], end=imported['report']['end'])
    response = client.post('/api/analytics/hazards', json=s)
    assert response.status_code == 200
    result = response.json()
    assert result['count'] == 2 and result['unlocated'] == 1
    assert all(e['group'] == 'Natural' for e in result['events'])
    assert len(result['regions']) == 27
    xlsx = client.post('/api/analytics/hazards/export/xlsx', json=s)
    assert xlsx.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content), read_only=True)
    assert book['Області'].max_row == 28
    assert book['Події'].max_row == 3
    book.close()
    csv = client.post('/api/analytics/hazards/export/csv', json=s)
    assert csv.status_code == 200 and 'Історичний індикатор' in csv.content.decode('utf-8-sig')
    assert client.post('/api/analytics/hazards', json={**s, 'frequency_weight': 0, 'severity_weight': 0}).status_code == 422
